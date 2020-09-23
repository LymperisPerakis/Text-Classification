import pandas as pd
from src.train_classification_model import TrainClassificationModel
import time
from typing import List

model = 'fasttext'
train_model = TrainClassificationModel(model=model)
documents_list, label_list, labels, label_dummy, documents = train_model.load_files()
fitted_model = train_model.train_model(test_size=0.2, random_state=0, epoch=10, lr=1, wordNgrams=3)
accuracy, report, confusion_matrix, predicted = train_model.validate_model()

from src.train_classification_model import TrainClassificationModel
from src.configs.configs import transformers_models

model = 'distilbert-base-uncased'
train_model = TrainClassificationModel(model=model)
documents_list, label_list, labels, label_dummy, documents = train_model.load_files()
X_train, X_test, y_train, y_test = train_model.split_training_files()
fitted_model = train_model.train_model()
accuracy, report, confusion_matrix, predicted = train_model.validate_model()

from src.train_classification_model import TrainClassificationModel
from src.configs.configs import conventional_models
import time
from typing import List
import pandas as pd
import docx
from random import randint
import seaborn as sns
sns.displot(label_dummy)

def find_average(metric: List[List[float]]):
    averages = []
    for x in metric:
        x_ = x if type(x) is list else [x]
        averages.append(sum(x_) / float(len(x_)))
    return averages


acc, pre, rec, f1_score, training_time, validation_time = [], [], [], [], [], []
random_state = [randint(0, 100) for i in range(5)]
for model in conventional_models:
    temp_acc, temp_pre, temp_rec, temp_f1, temp_training_time, temp_validation_time = [], [], [], [], [], []
    for i in range(5):
        train_model = TrainClassificationModel(model=model)
        documents_list, label_list, labels, label_dummy, documents = train_model.load_files()
        X_train, X_test, y_train, y_test = train_model.split_training_files(test_size=0.2, random_state=random_state[i])
        start = time.process_time()
        fitted_model = train_model.train_model()
        temp_training_time.append(time.process_time() - start)
        start = time.process_time()
        accuracy, report, confusion_matrix, predicted, precision, recall, f1 = train_model.validate_model()
        temp_validation_time.append(time.process_time() - start)
        temp_acc.append(accuracy)
        temp_pre.append(precision)
        temp_rec.append(recall)
        temp_f1.append(f1)
    acc.append(temp_acc)
    pre.append(temp_pre)
    rec.append(temp_rec)
    f1_score.append(temp_f1)
    training_time.append(temp_training_time)
    validation_time.append(temp_validation_time)
    print(f'Finished {model}')

data = {'Method': conventional_models,
        'Accuracy': find_average(acc),
        'Precision': find_average(pre),
        'Recall': find_average(rec),
        'F1': find_average(f1_score),
        }


df = pd.DataFrame(data, columns=['Method', 'Accuracy', 'Precision', 'Recall', 'F1'])
list = [round(x, 2) for x in list]

doc = docx.Document('Text.docx')
t = doc.add_table(df.shape[0]+1, df.shape[1])
t.style = 'TableGrid'
for j in range(df.shape[-1]):
    t.cell(0, j).text = df.columns[j]

# add the rest of the data frame
for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        t.cell(i+1, j).text = str(df.values[i, j])

# save the doc
doc.save('Text.docx')